"""P4: async jobs, worker retries, signed callbacks, export, review."""

import csv
import io
import json

import httpx
import pytest
from conftest import make_key
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core import signing
from app.models.db import JobReview, WebhookDelivery
from app.models.invoice import INVOICE_FIELDS
from app.services import callbacks
from app.services.llm import LLMError
from app.workers import tasks
from evals.generate_invoices import build

N8N_RESUME = "https://localhost:5678/webhook-waiting/12345"
FAKE_DNS = {"host.docker.internal": "192.168.65.254", "hooks.example-public.com": "93.184.216.34", "evil.internal": "10.0.0.7", "localhost": "127.0.0.1"}


def llm_out(data: dict, conf: float = 0.95) -> dict:
    return {**data, "field_confidence": {f: conf for f in INVOICE_FIELDS}}


@pytest.fixture(scope="module")
def sample():
    return build(1)


@pytest.fixture
def worker(monkeypatch, fake_llm):
    """Runs the Celery tasks in-process: records queued work, fakes the LLM, DNS and the callback receiver."""
    queued, callbacks_queued, received = [], [], []

    def receiver(request: httpx.Request) -> httpx.Response:
        received.append(request)
        return httpx.Response(worker.callback_status)

    monkeypatch.setattr(tasks, "enqueue_job", queued.append)
    monkeypatch.setattr(tasks, "enqueue_callback", callbacks_queued.append)
    monkeypatch.setattr(tasks, "get_llm", lambda: fake_llm)
    monkeypatch.setattr(callbacks, "_resolve", lambda host: [FAKE_DNS[host]])
    monkeypatch.setattr(callbacks, "TRANSPORT", httpx.MockTransport(receiver))
    worker.callback_status = 200
    worker.queued, worker.callbacks_queued, worker.received = queued, callbacks_queued, received
    return worker


def submit(client, key, content, *, callback_url=N8N_RESUME, idem=None, name="invoice.pdf"):
    headers = {"X-API-Key": key}
    if idem:
        headers["Idempotency-Key"] = idem
    data = {"schema": "invoice"}
    if callback_url:
        data["callback_url"] = callback_url
    return client.post("/v1/jobs", files={"file": (name, content, "application/pdf")}, data=data, headers=headers)


def run_job(job_id):
    return tasks.process_job.apply(args=[job_id])


# ---------- submission & idempotency ----------


def test_submit_returns_202_and_queues(client, api_key, worker, sample):
    r = submit(client, api_key, sample[0])
    assert r.status_code == 202
    body = r.json()
    assert body["status"] == "queued" and body["job_id"].startswith("job_")
    assert r.headers["Location"] == f"/v1/jobs/{body['job_id']}"
    assert worker.queued == [body["job_id"]]
    job = client.get(f"/v1/jobs/{body['job_id']}", headers={"X-API-Key": api_key}).json()
    assert job["status"] == "queued" and job["data"] is None


def test_same_file_returns_same_job_without_a_key(client, api_key, worker, sample):
    first = submit(client, api_key, sample[0])
    second = submit(client, api_key, sample[0])
    assert second.json()["job_id"] == first.json()["job_id"]
    assert second.headers["Idempotent-Replay"] == "true"
    assert len(worker.queued) == 1


def test_explicit_idempotency_key_is_honoured(client, api_key, worker, sample):
    a = submit(client, api_key, sample[0], idem="intake-42")
    b = submit(client, api_key, sample[0], idem="intake-42")
    assert a.json()["job_id"] == b.json()["job_id"]
    assert submit(client, api_key, build(2)[0], idem="intake-42").status_code == 409  # same key, different file


@pytest.mark.parametrize(
    "url", ["http://evil.internal/hook", "https://localhost:9999/x", "ftp://hooks.example-public.com/x", "https://user:pw@hooks.example-public.com/"]
)
def test_unsafe_callback_urls_are_rejected(client, api_key, worker, sample, url):
    r = submit(client, api_key, sample[0], callback_url=url)
    assert r.status_code == 422 and "callback_url not allowed" in r.json()["detail"]
    assert worker.queued == []


def test_non_pdf_is_415(client, api_key, worker):
    assert submit(client, api_key, b"hello world").status_code == 415


# ---------- worker ----------


def test_worker_extracts_and_sends_signed_callback(client, api_key, worker, fake_llm, sample, engine):
    content, label = sample
    job_id = submit(client, api_key, content).json()["job_id"]
    fake_llm.queue = [llm_out(label["data"])]
    assert run_job(job_id).get() == "done"
    assert worker.callbacks_queued == [job_id]

    assert tasks.send_callback.apply(args=[job_id]).get() == "delivered"
    sent = worker.received[0]
    # connection goes to the mapped target IP, while Host + TLS SNI keep the name the client gave
    assert sent.url.host == "192.168.65.254" and sent.url.port == 5678
    assert sent.headers["host"] == "localhost:5678" and sent.extensions["sni_hostname"] == "localhost"
    assert signing.verify("test-callback-secret", sent.headers[signing.TIMESTAMP_HEADER], sent.content, sent.headers[signing.SIGNATURE_HEADER])
    assert not signing.verify("test-callback-secret", sent.headers[signing.TIMESTAMP_HEADER], sent.content + b" ", sent.headers[signing.SIGNATURE_HEADER])

    job = client.get(f"/v1/jobs/{job_id}", headers={"X-API-Key": api_key}).json()
    assert json.loads(sent.content) == job
    assert job["status"] == "done" and job["data"] == label["data"]
    with Session(engine) as s:
        rows = s.scalars(select(WebhookDelivery).where(WebhookDelivery.job_id == job_id)).all()
    assert [(r.attempt, r.status_code) for r in rows] == [(1, 200)]


def test_retryable_llm_error_is_retried_then_succeeds(client, api_key, worker, fake_llm, sample, engine):
    content, label = sample
    job_id = submit(client, api_key, content).json()["job_id"]
    fake_llm.queue = [LLMError("connection error", retryable=True), llm_out(label["data"])]
    run_job(job_id)
    job = client.get(f"/v1/jobs/{job_id}", headers={"X-API-Key": api_key}).json()
    assert job["status"] == "done"
    with engine.connect() as conn:
        attempts, last_error = conn.exec_driver_sql("SELECT attempts, last_error FROM jobs WHERE id = %s", (job_id,)).one()
    assert attempts == 2 and "attempt 1" in last_error
    assert worker.callbacks_queued == [job_id]


def test_gives_up_after_three_attempts_and_still_calls_back(client, api_key, worker, fake_llm, sample, engine):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    fake_llm.queue = [LLMError("timeout", retryable=True)] * 3 + [llm_out(sample[1]["data"])]
    run_job(job_id)
    job = client.get(f"/v1/jobs/{job_id}", headers={"X-API-Key": api_key}).json()
    assert job["status"] == "failed" and job["validation"]["errors"][0]["code"] == "llm_unavailable"
    with engine.connect() as conn:
        assert conn.exec_driver_sql("SELECT attempts FROM jobs WHERE id = %s", (job_id,)).scalar() == 3
    assert fake_llm.queue  # the 4th response was never requested
    assert worker.callbacks_queued == [job_id]


def test_non_retryable_llm_error_fails_at_once(client, api_key, worker, fake_llm, sample, engine):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    fake_llm.queue = [LLMError("provider returned 400")]
    run_job(job_id)
    with engine.connect() as conn:
        status, attempts = conn.exec_driver_sql("SELECT status, attempts FROM jobs WHERE id = %s", (job_id,)).one()
    assert (status, attempts) == ("failed", 1)


def test_failed_callbacks_are_logged_and_retried_then_given_up(client, api_key, worker, fake_llm, sample, engine):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    fake_llm.queue = [llm_out(sample[1]["data"])]
    run_job(job_id)
    worker.callback_status = 500
    tasks.send_callback.apply(args=[job_id])
    with Session(engine) as s:
        rows = s.scalars(select(WebhookDelivery).where(WebhookDelivery.job_id == job_id).order_by(WebhookDelivery.attempt)).all()
    assert [r.attempt for r in rows] == [1, 2, 3, 4, 5, 6]
    assert all(r.status_code == 500 for r in rows)


@pytest.mark.parametrize("status,expected_attempts", [(409, 1), (404, 1), (429, 6), (408, 6)])
def test_permanent_4xx_callbacks_are_not_retried(client, api_key, worker, fake_llm, sample, engine, status, expected_attempts):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    fake_llm.queue = [llm_out(sample[1]["data"])]
    run_job(job_id)
    worker.callback_status = status
    tasks.send_callback.apply(args=[job_id])
    with Session(engine) as s:
        rows = s.scalars(select(WebhookDelivery).where(WebhookDelivery.job_id == job_id)).all()
    assert len(rows) == expected_attempts
    assert all(r.status_code == status for r in rows)


def test_duplicate_detected_across_async_jobs(client, api_key, worker, fake_llm, sample):
    content, label = sample
    first = submit(client, api_key, content, idem="a").json()["job_id"]
    second = submit(client, api_key, content, idem="b").json()["job_id"]
    fake_llm.queue = [llm_out(label["data"]), llm_out(label["data"])]
    run_job(first)
    run_job(second)
    job = client.get(f"/v1/jobs/{second}", headers={"X-API-Key": api_key}).json()
    assert job["duplicate_of"] == first


# ---------- read, export, review ----------


@pytest.fixture
def finished_job(client, api_key, worker, fake_llm, sample):
    content, label = sample
    job_id = submit(client, api_key, content).json()["job_id"]
    fake_llm.queue = [llm_out(label["data"])]
    run_job(job_id)
    return job_id, label


def test_export_csv_and_xlsx(client, api_key, finished_job):
    job_id, label = finished_job
    r = client.get(f"/v1/jobs/{job_id}/export?format=csv", headers={"X-API-Key": api_key})
    assert r.status_code == 200 and r.headers["content-type"].startswith("text/csv")
    rows = list(csv.DictReader(io.StringIO(r.content.decode("utf-8-sig"))))
    assert len(rows) == len(label["data"]["line_items"])
    assert rows[0]["invoice_no"] == label["data"]["invoice_no"] and rows[0]["line_description"] == label["data"]["line_items"][0]["description"]

    x = client.get(f"/v1/jobs/{job_id}/export?format=xlsx", headers={"X-API-Key": api_key})
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.sheetnames == ["Invoice", "Line items"]
    assert dict(wb["Invoice"].iter_rows(values_only=True))["total"] == label["data"]["total"]
    assert wb["Line items"].max_row == len(label["data"]["line_items"]) + 1


def test_export_before_extraction_is_409(client, api_key, worker, sample):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    assert client.get(f"/v1/jobs/{job_id}/export", headers={"X-API-Key": api_key}).status_code == 409


def test_review_stores_correction_separately(client, api_key, finished_job, engine):
    job_id, label = finished_job
    corrected = {**label["data"], "vendor": "Bluefin Hosting Limited"}
    r = client.post(f"/v1/jobs/{job_id}/review", json={"data": corrected, "notes": "legal name"}, headers={"X-API-Key": api_key})
    assert r.status_code == 200 and r.json()["status"] == "reviewed" and r.json()["validation"]["ok"] is True
    job = client.get(f"/v1/jobs/{job_id}", headers={"X-API-Key": api_key}).json()
    assert job["status"] == "reviewed" and job["data"]["vendor"] == label["data"]["vendor"]  # model output untouched
    with Session(engine) as s:
        review = s.scalars(select(JobReview).where(JobReview.job_id == job_id)).one()
    assert review.data["vendor"] == "Bluefin Hosting Limited" and review.notes == "legal name"


def test_review_reports_validation_errors(client, api_key, finished_job):
    job_id, label = finished_job
    r = client.post(f"/v1/jobs/{job_id}/review", json={"data": {**label["data"], "total": 1.0}}, headers={"X-API-Key": api_key})
    assert r.json()["validation"]["ok"] is False and "total_mismatch" in {e["code"] for e in r.json()["validation"]["errors"]}


def test_review_of_unfinished_job_is_409(client, api_key, worker, sample):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    r = client.post(f"/v1/jobs/{job_id}/review", json={"data": sample[1]["data"]}, headers={"X-API-Key": api_key})
    assert r.status_code == 409


def test_list_filters_by_status_and_is_scoped_to_the_client(client, engine, api_key, worker, fake_llm, sample):
    other = make_key(engine, name="other-client")
    submit(client, api_key, sample[0])
    submit(client, other, build(2)[0])
    mine = client.get("/v1/jobs?status=queued", headers={"X-API-Key": api_key}).json()
    assert len(mine) == 1 and mine[0]["status"] == "queued"
    assert client.get("/v1/jobs?status=done", headers={"X-API-Key": api_key}).json() == []


def test_other_clients_job_is_404(client, engine, api_key, worker, sample):
    job_id = submit(client, api_key, sample[0]).json()["job_id"]
    other = make_key(engine, name="other-client")
    assert client.get(f"/v1/jobs/{job_id}", headers={"X-API-Key": other}).status_code == 404


def test_job_backoff_has_a_floor_and_grows():
    for _ in range(200):
        assert 5 <= tasks.job_backoff_s(1) <= 10
        assert 10 <= tasks.job_backoff_s(2) <= 20
