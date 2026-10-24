"""Spreadsheet export of one extracted invoice: header fields plus line items."""

import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font

from app.models.invoice import ExtractResponse

HEADER_FIELDS = ("vendor", "invoice_no", "invoice_date", "due_date", "currency", "subtotal", "tax", "total")
LINE_FIELDS = ("description", "qty", "unit_price", "amount")


def _header_row(resp: ExtractResponse) -> dict:
    data = resp.data.model_dump(mode="json")
    return {
        "job_id": resp.job_id,
        "status": resp.status,
        **{f: data[f] for f in HEADER_FIELDS},
        "confidence": resp.confidence.overall,
        "validation_ok": resp.validation.ok,
    }


def to_csv(resp: ExtractResponse) -> bytes:
    """One row per line item with the invoice header repeated, so it pastes straight into a ledger."""
    header = _header_row(resp)
    columns = [*header, *(f"line_{f}" for f in LINE_FIELDS)]
    buf = io.StringIO()
    writer = csv.DictWriter(buf, fieldnames=columns)
    writer.writeheader()
    items = resp.data.line_items or [None]
    for item in items:
        row = dict(header)
        if item is not None:
            row.update({f"line_{f}": getattr(item, f) for f in LINE_FIELDS})
        writer.writerow(row)
    return buf.getvalue().encode("utf-8-sig")  # BOM so Excel opens UTF-8 correctly


def to_xlsx(resp: ExtractResponse) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoice"
    for key, value in _header_row(resp).items():
        ws.append([key, value])
    for cell in ws["A"]:
        cell.font = Font(bold=True)
    lines = wb.create_sheet("Line items")
    lines.append(list(LINE_FIELDS))
    for cell in lines[1]:
        cell.font = Font(bold=True)
    for item in resp.data.line_items:
        lines.append([getattr(item, f) for f in LINE_FIELDS])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
